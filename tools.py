"""The calculation tools exposed to the Strands agent.

All arithmetic happens here (Python); the AI never does mental math.
Structure-changing tools verify the passphrase and log before/after
values to settings_log. Every string the owner sees comes from i18n.py
(``t("key")``), so the same code speaks English or Japanese.
"""

from __future__ import annotations

import csv
import functools
import re
from contextvars import ContextVar
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from strands import tool

from i18n import t
from store import (
    STATE_LOCK,
    STATE_PATH,
    StateConflictError,
    archive_and_reset,
    load_state,
    save_state,
)


def _serialized(func):
    """Serialize one whole load→save cycle per tool call.

    The agent sometimes calls tools in parallel (e.g. during stock counts);
    without the lock a later write silently clobbers an earlier one
    (observed on real hardware).
    """

    @functools.wraps(func)
    def wrapper(*args, **kwargs):
        with STATE_LOCK:
            return func(*args, **kwargs)

    return wrapper


TIMEZONE = ZoneInfo("Asia/Tokyo")
# Monday-first, to match datetime.weekday(). Named rather than taken from the
# C locale, which is not installed in every container teruo runs in.
WEEKDAYS = (
    "weekday_mon", "weekday_tue", "weekday_wed", "weekday_thu",
    "weekday_fri", "weekday_sat", "weekday_sun",
)
VENUE_TYPES = ("event", "solo")
UNIT_TYPES = ("weight", "volume", "count")
# The three consumption types (design doc §2). This is the counting axis,
# separate from unit_type (the kind of unit).
CONSUMPTION_TYPES = ("count", "weight", "unit")
# Conflict / passphrase / negative-stock wording: see i18n.py
# (t("conflict"), t("passphrase_wrong"), t("negative_stock")).

# Conversion tables within the same unit kind (multiplier per base unit).
# Never convert across kinds. Metric and imperial both live here — which
# system a shop uses is decided at onboarding, per item.
WEIGHT_UNITS = {"g": 1.0, "kg": 1000.0, "oz": 28.3495, "lb": 453.59237}
VOLUME_UNITS = {"ml": 1.0, "mL": 1.0, "l": 1000.0, "L": 1000.0, "fl oz": 29.5735}
COUNT_UNITS = (
    "pc", "sheet", "bottle", "bag", "can", "box",
    "serving", "cup", "roll", "tub", "piece", "cone",
    # Japanese counters — a shop registers whichever it actually counts in
    "枚", "本", "個", "袋", "缶", "箱", "食", "杯", "巻", "中子",
)

ITEM_ID_PATTERN = re.compile(r"^[a-z][a-z0-9_]*$")

# Learning-phase cutoff: 5 stock counts or 14 days (design doc, principle 9)
LEARNING_MAX_COUNTS = 5
LEARNING_MAX_DAYS = 14
STABLE_COEFFICIENT_BAND = 0.05

# Coefficient cap (design doc §2). Portioning variance is physically
# bounded at about ±3-4g per serving. Defined in grams; always convert to
# the item's unit via _bound_per_serving (±4 would mean ±113g for an oz item).
WEIGHT_BOUND_PER_SERVING = 4.0
# Coefficient smoothing. Adopting a single count's correction in full chases
# day-to-day hand variance (about ±5%) and oscillates (confirmed by the
# convergence simulation). Moving halfway keeps ~5-count convergence while
# halving the swing.
COEFFICIENT_SMOOTHING = 0.5
# Unit-tracked items: ±20% of past records. No cap check until 3 records exist.
UNIT_BOUND_RATIO = 0.2
UNIT_BOUND_MIN_SAMPLES = 3
UNIT_LEARNING_SAMPLES = 3


def _now() -> datetime:
    return datetime.now(TIMEZONE)


def today() -> str:
    """The current date in the shop's timezone as YYYY-MM-DD. The model has
    no clock; the prompts get the date from here."""
    return _now().strftime("%Y-%m-%d")


def now_stamp() -> str:
    """The current date and time in the shop's timezone, spelled for the owner.

    Both entry points lead with this. It is the one line that shows the clock
    the records are stamped from is the clock the owner is standing in — a
    machine an hour off, or a day behind, is worth catching before the first
    sale is recorded rather than at the month's reconciliation.
    """
    moment = _now()
    return t(
        "clock",
        date=moment.strftime("%Y-%m-%d"),
        weekday=t(WEEKDAYS[moment.weekday()]),
        time=moment.strftime("%H:%M"),
    )


def intro_line() -> str:
    """The first line both entry points print: whose shop this is, and the clock.

    The shop's name once onboarding has captured it; the generic description
    before that, because on a first run teruo genuinely does not know whose
    shop it is and should not pretend otherwise. Read defensively — this runs
    before anything else, and a missing or half-written state file should not
    stop the owner from reaching the prompt.
    """
    try:
        shop = (load_state().get("config") or {}).get("shop_name")
    except (OSError, ValueError):
        shop = None
    if shop:
        return t("intro_named", shop=shop, clock=now_stamp())
    return t("intro", clock=now_stamp())


def _now_iso() -> str:
    """Timestamps are always generated in Python. Never let the AI invent dates."""
    return _now().isoformat()


def _parse_iso(value: str) -> datetime:
    return datetime.fromisoformat(value)


def _fmt_dt(value: str) -> str:
    try:
        return _parse_iso(value).strftime("%Y-%m-%d %H:%M")
    except ValueError:
        return value


def _round_one(value: float) -> float:
    return float(Decimal(str(value)).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP))


def _display_number(value: float) -> str:
    return str(int(value)) if float(value).is_integer() else f"{value:.1f}"


def _plural(unit: str, value: float) -> str:
    if abs(value) == 1 or not unit.isascii():
        return unit
    if unit.endswith(("s", "x", "ch", "sh")):
        return unit + "es"
    return unit + "s"


def _amount(value: float, unit: str) -> str:
    """Format a quantity with its unit: '300g' for measures, '3 bottles' for
    counts. Japanese counters neither pluralize nor take a space: '3本'."""
    if unit in WEIGHT_UNITS or unit in VOLUME_UNITS or not unit.isascii():
        return f"{_display_number(value)}{unit}"
    return f"{_display_number(value)} {_plural(unit, value)}"


def _join(names) -> str:
    """Join display names with the language's list separator (', ' or '、')."""
    return t("list_sep").join(names)


def _bound_per_serving(unit: str) -> float:
    """The portioning-variance cap expressed in the item's unit.

    The physical limit is defined in grams (±4g/serving); ml is treated as
    the equivalent for volume items. Unknown units fall back to the raw value.
    """
    factor = WEIGHT_UNITS.get(unit) or VOLUME_UNITS.get(unit) or 1.0
    return WEIGHT_BOUND_PER_SERVING / factor


def _display_bound(value: float) -> str:
    """Bound values below 1 (e.g. 0.14oz) need two decimals to stay meaningful."""
    return _display_number(value) if value >= 1 else f"{value:.2f}"


# Facts (final tool output) are printed straight to the screen by Python
# (principle 3: Python calculates, AI judges). This prevents the LLM's
# paraphrase from dropping line items or altering numbers.
# seed_demo.py and friends set this to False to silence it.
DIRECT_OUTPUT = True

# Where a fact goes once Python has produced it. The CLI leaves this unset and
# the text lands on stdout; the web entry point installs a sink so the same
# text reaches that browser instead. A ContextVar rather than a global because
# Strands runs sync tools through asyncio.to_thread, which copies the context
# into the worker thread — so the fact follows the request it belongs to.
_OUTPUT_SINK: ContextVar[Any] = ContextVar("teruo_output_sink", default=None)


def set_output_sink(sink) -> Any:
    """Route facts to ``sink`` instead of stdout. Returns a reset token."""
    return _OUTPUT_SINK.set(sink)


def reset_output_sink(token: Any) -> None:
    _OUTPUT_SINK.reset(token)


# The same idea for files. An export writes real files somewhere on disk; at
# the CLI the path it prints is enough, but a browser cannot reach a path. The
# web entry point installs a sink here and turns what it receives into a
# download link, so "hand it to me" works from either entry point.
_FILE_SINK: ContextVar[Any] = ContextVar("teruo_file_sink", default=None)


def set_file_sink(sink) -> Any:
    """Route written files to ``sink``. Returns a reset token."""
    return _FILE_SINK.set(sink)


def reset_file_sink(token: Any) -> None:
    _FILE_SINK.reset(token)


def _offer(paths: list[Path]) -> None:
    """Announce files just written. No sink (the CLI) means nothing happens —
    the tool's own text already says where they are."""
    sink = _FILE_SINK.get()
    if sink is not None:
        sink(list(paths))


def told_marker() -> str:
    """The tag telling the LLM a tool result is already on screen (per language)."""
    return t("told_marker")


def hit_cap(result: str) -> bool:
    """Whether a record_count result carries the coefficient-cap warning."""
    return t("cap_hit_fragment") in result


def _tell(text: str) -> str:
    """Emit the final output to the owner's screen and tag it so the LLM won't echo it."""
    if DIRECT_OUTPUT:
        sink = _OUTPUT_SINK.get()
        if sink is None:
            print(text, flush=True)
        else:
            sink(text)
        return f"{told_marker()}\n{text}"
    return text


def _known_ids(records: list[dict[str, Any]]) -> str:
    ids = [r["id"] for r in records if r.get("active", True)]
    return ", ".join(ids) if ids else t("known_ids_none")


def _item_missing(state: dict[str, Any], key: str, item_id: str) -> str:
    """An item-ID error plus the registered IDs, so a typo is fixed in one
    step instead of guessed at."""
    return f'{t(key, item_id=item_id)}\n{t("known_item_ids", ids=_known_ids(state["items"]))}'


def _product_missing(state: dict[str, Any], product_id: str) -> str:
    return (
        f'{t("product_not_found", product_id=product_id)}\n'
        f'{t("known_product_ids", ids=_known_ids(state["products"]))}'
    )


def _find(records: list[dict[str, Any]], record_id: str) -> dict[str, Any] | None:
    return next((record for record in records if record["id"] == record_id), None)


def _consumption_type(item: dict[str, Any]) -> str:
    """Consumption type. Older data without one falls back to unit_type
    (count→count, everything else→weight)."""
    value = item.get("consumption_type")
    if value in CONSUMPTION_TYPES:
        return value
    return "count" if item.get("unit_type") == "count" else "weight"


def _display_stock(item: dict[str, Any]) -> str:
    """Stock for display. Negative values stay stored internally but the
    number itself is never shown (design doc, principle 10)."""
    stock = float(item["stock"])
    if stock < 0:
        return t("negative_stock")
    return _amount(stock, item["unit"])


def _check_passphrase(state: dict[str, Any], passphrase: str) -> str | None:
    """Match against the stored passphrase. Pass through when none is set
    (i.e. during onboarding)."""
    stored = (state.get("config") or {}).get("passphrase")
    if not stored:
        return None
    if passphrase == stored:
        return None
    return t("passphrase_wrong")


def _log_setting(
    state: dict[str, Any],
    tool_name: str,
    summary: str,
    before: Any,
    after: Any,
) -> None:
    """Structure-change history. Always keep both the before and the after
    (design doc, principle 4)."""
    state.setdefault("settings_log", []).append(
        {
            "changed_at": _now_iso(),
            "tool": tool_name,
            "summary": summary,
            "before": before,
            "after": after,
        }
    )


def _unit_kind(unit: str) -> str | None:
    if unit in WEIGHT_UNITS:
        return "weight"
    if unit in VOLUME_UNITS:
        return "volume"
    if unit in COUNT_UNITS:
        return "count"
    return None


def _conversion_factor(old_unit: str, new_unit: str) -> float | None:
    """Conversion multiplier within the same kind (g→kg etc.). None across kinds."""
    for table in (WEIGHT_UNITS, VOLUME_UNITS):
        if old_unit in table and new_unit in table:
            return table[old_unit] / table[new_unit]
    return None


def _growth_counts(item: dict[str, Any]) -> list[dict[str, Any]]:
    """Stock counts since the most recent unit change. A unit change
    restarts the learning phase."""
    counts: list[dict[str, Any]] = []
    for entry in item.get("count_history", []):
        if entry.get("event") == "unit_change":
            counts = []
        elif "actual_stock" in entry:
            counts.append(entry)
    return counts


def _growth(item: dict[str, Any]) -> tuple[str, str]:
    """Learning phase as (code, label) (design doc, principle 9).

    Codes: fixed / unlearned / learning / stable / insufficient / unsettled.
    Callers branch on the code, never on the label's wording.
    "learning" is cut off unconditionally after 5 stock counts or 14 days,
    and never said again. If the coefficient hasn't settled, return the fact
    (recent coefficient movement) and leave interpretation to the AI.
    Unit-tracked items count empty-unit records instead (design doc §2).
    """
    ctype = _consumption_type(item)
    if ctype == "count":
        return "fixed", t("growth_fixed")
    if ctype == "unit":
        n = len(item.get("unit_history", []))
        if n == 0:
            return "unlearned", t("growth_unlearned")
        if n < UNIT_LEARNING_SAMPLES:
            return "learning", t("growth_learning_unit", n=n, max=UNIT_LEARNING_SAMPLES)
        return "stable", t("growth_stable_unit", n=n)
    counts = _growth_counts(item)
    n = len(counts)
    if n == 0:
        return "learning", t("growth_learning", n=0, max=LEARNING_MAX_COUNTS)
    days = (_now() - _parse_iso(counts[0]["recorded_at"])).days
    if n < LEARNING_MAX_COUNTS and days < LEARNING_MAX_DAYS:
        return "learning", t("growth_learning", n=n, max=LEARNING_MAX_COUNTS)
    if n < 2:
        return "insufficient", t("growth_insufficient")
    delta = abs(
        float(counts[-1]["coefficient_after"]) - float(counts[-2]["coefficient_after"])
    )
    if delta <= STABLE_COEFFICIENT_BAND:
        return "stable", t("growth_stable")
    return "unsettled", t(
        "growth_unsettled",
        before=float(counts[-2]["coefficient_after"]),
        after=float(counts[-1]["coefficient_after"]),
    )


def _growth_label(item: dict[str, Any]) -> str:
    return _growth(item)[1]


def _save(state: dict[str, Any]) -> str | None:
    try:
        save_state(state)
    except StateConflictError:
        return t("conflict")
    return None


# ---------------------------------------------------------------------------
# Daily entries (no passphrase required)
# ---------------------------------------------------------------------------


@tool
@_serialized
def record_sales(product_id: str, quantity: int, venue_type: str = "solo") -> str:
    """Record sales and reduce item stock according to the product's recipe.

    Each consumption type behaves differently (design doc §2):
    count items decrease one-for-one, weight items by recipe × coefficient.
    Unit items (sauce etc.) only accumulate a serving tally; their stock
    moves when an empty unit is recorded.
    The sale timestamp is set by Python automatically — no date input needed.

    Args:
        product_id: Product ID, e.g. kebab_sand (kebab sandwich).
        quantity: Number sold. Integer, zero or more.
        venue_type: "event" (event booth) or "solo" (regular solo pitch).
            Use "event" only when the owner explicitly says it was an event.
            Defaults to "solo".
    """
    if quantity < 0:
        return t("qty_nonnegative")
    if venue_type not in VENUE_TYPES:
        return t("venue_invalid")

    state = load_state()
    product = _find(state["products"], product_id)
    if product is None:
        return _product_missing(state, product_id)

    changes: list[str] = []
    notes: list[str] = []
    for ingredient in product["recipe"]:
        item = _find(state["items"], ingredient["item_id"])
        if item is None or not item.get("active", True):
            continue
        ctype = _consumption_type(item)
        qty = float(ingredient["qty"])

        if ctype == "unit":
            # Partial contents can't be measured. Only tally servings
            # (stock moves via record_unit_used).
            item["sales_count"] = float(item.get("sales_count", 0)) + quantity * qty
            coefficient = item.get("coefficient")
            opened_at = item.get("opened_at_sales_count")
            used_note = ""
            if opened_at is not None:
                used = int(round(float(item["sales_count"]) - float(opened_at)))
                used_note = t("sales_open_note", unit=item["unit"], used=used)
            changes.append(
                t(
                    "sales_unit_unchanged",
                    name=item["name"],
                    total=_display_number(float(item["sales_count"])),
                    used_note=used_note,
                )
            )
            if coefficient and opened_at is not None:
                remaining = float(coefficient) - used
                if remaining <= max(5.0, float(coefficient) * 0.1):
                    notes.append(
                        t(
                            "sales_unit_running_out",
                            name=item["name"],
                            unit=item["unit"],
                            remaining=max(0, int(remaining)),
                        )
                    )
            continue

        coefficient = float(item["coefficient"]) if ctype == "weight" else 1.0
        before = float(item["stock"])
        consumed = _round_one(quantity * qty * coefficient)
        after = _round_one(before - consumed)
        item["stock"] = after
        if after < 0 <= before:
            # Keep the negative value, never show the number (principle 10).
            # Don't stop with an error either.
            changes.append(t("sales_went_negative", name=item["name"]))
        elif after < 0:
            changes.append(f'{item["name"]}: {t("negative_stock")}')
        else:
            changes.append(
                t(
                    "sales_decrease",
                    name=item["name"],
                    before=_display_number(before),
                    after=_amount(after, item["unit"]),
                    consumed=_amount(consumed, item["unit"]),
                )
            )

    state["history"].append(
        {
            "recorded_at": _now_iso(),
            "venue_type": venue_type,
            "product_id": product_id,
            "quantity": quantity,
        }
    )
    if conflict := _save(state):
        return conflict
    lines = [t("sales_recorded", product=product["name"], quantity=quantity)]
    lines.extend(f"- {change}" for change in changes)
    lines.extend(notes)
    return _tell("\n".join(lines))


@tool
@_serialized
def record_count(item_id: str, actual_stock: float) -> str:
    """Record a physical stock count. For weight items, the difference since
    the last count corrects the consumption coefficient.

    Count items are only recounted (coefficient fixed at 1.0). Unit items are
    only recounted in units; their coefficient is settled by empty-unit
    records (record_unit_used).

    Args:
        item_id: Item ID, e.g. meat_chicken (kebab meat, chicken).
        actual_stock: The counted stock, in the item's registered unit.
            For unit items, count an opened unit as 1.
    """
    if actual_stock < 0:
        return t("count_nonnegative")

    state = load_state()
    item = _find(state["items"], item_id)
    if item is None or not item.get("active", True):
        return _item_missing(state, "item_not_found", item_id)

    ctype = _consumption_type(item)
    counted_at = _now_iso()
    previous_coefficient = float(item.get("coefficient") or 1.0)
    calculated_stock = float(item["stock"])
    last_counted = item.get("last_counted")

    warning = ""
    message: str
    if ctype == "unit":
        message = t("count_unit_recounted", name=item["name"], unit=item["unit"])
    elif ctype == "count":
        difference = _round_one(calculated_stock - actual_stock)
        if difference == 0:
            message = t("count_matches", name=item["name"])
        else:
            # A gap in a count-tracked item can't be explained by a coefficient.
            # State the fact only (principle 12).
            direction = t("fewer") if difference > 0 else t("more")
            message = t(
                "count_gap",
                name=item["name"],
                diff=_amount(abs(difference), item["unit"]),
                direction=direction,
            )
    elif last_counted is None:
        message = t("count_first", name=item["name"], coef=previous_coefficient)
    else:
        last_counted_at = _parse_iso(last_counted)
        theoretical_consumption = 0.0
        recipe_consumption = 0.0
        servings = 0.0
        for sale in state["history"]:
            recorded_at = sale.get("recorded_at")
            product_id = sale.get("product_id")
            if (
                recorded_at is None
                or product_id is None
                or _parse_iso(recorded_at) <= last_counted_at
            ):
                continue
            product = _find(state["products"], product_id)
            if product is None:
                continue
            recipe = next(
                (
                    ingredient
                    for ingredient in product["recipe"]
                    if ingredient["item_id"] == item_id
                ),
                None,
            )
            if recipe:
                quantity = float(sale["quantity"])
                theoretical_consumption += (
                    quantity * float(recipe["qty"]) * previous_coefficient
                )
                recipe_consumption += quantity * float(recipe["qty"])
                servings += quantity

        if theoretical_consumption > 0:
            actual_consumption = theoretical_consumption + (
                calculated_stock - actual_stock
            )
            measured_coefficient = previous_coefficient * (
                actual_consumption / theoretical_consumption
            )
            # Smoothing: never adopt one count in full, move halfway
            # (anti-oscillation).
            new_coefficient = previous_coefficient + COEFFICIENT_SMOOTHING * (
                measured_coefficient - previous_coefficient
            )
            # Coefficient cap: recipe value ±4g/serving (design doc §2) —
            # the physical limit of portioning variance, in the item's unit.
            bound = _bound_per_serving(item["unit"])
            base_per_serving = recipe_consumption / servings
            lower = max(0.0, base_per_serving - bound) / base_per_serving
            upper = (base_per_serving + bound) / base_per_serving
            clamped = min(upper, max(lower, new_coefficient))
            item["coefficient"] = round(clamped, 4)
            message = t(
                "count_coef_updated",
                name=item["name"],
                before=previous_coefficient,
                after=item["coefficient"],
            )
            if clamped != new_coefficient:
                warning = t(
                    "count_cap_hit", bound=_display_bound(bound), unit=item["unit"]
                )
            elif abs(clamped - previous_coefficient) > 0.15:
                warning = t("count_big_jump")
        else:
            message = t("count_no_sales", name=item["name"], coef=previous_coefficient)

    phase_before = _growth(item)[0]
    item["stock"] = _round_one(actual_stock)
    item["last_counted"] = counted_at
    item.setdefault("count_history", []).append(
        {
            "recorded_at": counted_at,
            "actual_stock": item["stock"],
            "calculated_stock": calculated_stock,
            "coefficient_before": previous_coefficient,
            "coefficient_after": float(item.get("coefficient") or 1.0),
        }
    )
    phase, growth = _growth(item)
    growth_note = ""
    if ctype == "weight" and phase_before == "learning" and phase != "learning":
        # When learning ends, always show the gap between the initial and the
        # settled value. Never finalize silently (design doc §2).
        recipe_values = [
            float(ingredient["qty"])
            for product in state["products"]
            for ingredient in product["recipe"]
            if ingredient["item_id"] == item_id
        ]
        if recipe_values:
            base = sum(recipe_values) / len(recipe_values)
            final = base * float(item["coefficient"])
            percent = (final - base) / base * 100
            growth_note = t(
                "count_learning_done",
                name=item["name"],
                base=_amount(_round_one(base), item["unit"]),
                final=_amount(_round_one(final), item["unit"]),
                percent=percent,
            )
    elif phase == "learning":
        growth_note = t("count_still_learning", growth=growth)
    elif phase == "unsettled":
        growth_note = t("count_unsettled_note", growth=growth)
    if conflict := _save(state):
        return conflict
    return _tell(
        t(
            "count_stock_updated",
            message=message,
            stock=_amount(float(item["stock"]), item["unit"]),
            warning=warning,
            growth_note=growth_note,
        ).rstrip()
    )


@tool
@_serialized
def record_unit_used(item_id: str, opened_next: bool = True) -> str:
    """Record that a unit-tracked item (sauce, oil, onion, ...) was used up.

    The serving tally since it was opened settles "servings per unit"
    (design doc §2). Partial contents are never handled — only the certain
    fact that a unit is empty feeds learning. If no opening point exists yet,
    this only records the opening point of the unit currently in use.

    Args:
        item_id: Item ID, e.g. sauce_yogurt.
        opened_next: Whether the next unit was opened right away after this
            one emptied. False if not.
    """
    state = load_state()
    item = _find(state["items"], item_id)
    if item is None or not item.get("active", True):
        return _item_missing(state, "item_not_found", item_id)
    if _consumption_type(item) != "unit":
        return t("unit_not_unit_type", name=item["name"])

    unit = item["unit"]
    sales_count = float(item.get("sales_count", 0))
    opened_at = item.get("opened_at_sales_count")

    if opened_at is None:
        item["opened_at_sales_count"] = sales_count
        if conflict := _save(state):
            return conflict
        return _tell(t("unit_opened", name=item["name"], unit=unit))

    servings = int(round(sales_count - float(opened_at)))
    history = item.setdefault("unit_history", [])
    anomaly = ""
    if servings <= 0:
        anomaly = t("unit_no_sales", unit=unit)
    elif len(history) >= UNIT_BOUND_MIN_SAMPLES:
        # Cap check: ±20% of past records (design doc §2). An out-of-band
        # value is kept as a fact but not learned from.
        average = sum(history) / len(history)
        if abs(servings - average) > UNIT_BOUND_RATIO * average:
            anomaly = t(
                "unit_out_of_band",
                servings=servings,
                unit=unit,
                pct=int(UNIT_BOUND_RATIO * 100),
                average=average,
            )
    if not anomaly:
        history.append(servings)
        item["coefficient"] = round(sum(history) / len(history), 1)

    item.setdefault("count_history", []).append(
        {
            "recorded_at": _now_iso(),
            "event": "unit_used",
            "servings": servings,
            "reflected": not anomaly,
            "coefficient_after": item.get("coefficient"),
        }
    )
    before = float(item["stock"])
    item["stock"] = _round_one(before - 1)
    item["opened_at_sales_count"] = sales_count if opened_next else None

    if conflict := _save(state):
        return conflict

    stock_display = _display_stock(item)
    if anomaly:
        return _tell(
            t("unit_anomaly_result", name=item["name"], anomaly=anomaly, stock=stock_display)
        )
    coefficient = float(item["coefficient"])
    capacity_note = (
        t("unit_capacity_note", servings=int(max(0.0, float(item["stock"])) * coefficient))
        if float(item["stock"]) >= 0
        else ""
    )
    return _tell(
        t(
            "unit_used_result",
            name=item["name"],
            servings=servings,
            unit=unit,
            count=len(history),
            coef=coefficient,
            stock=stock_display,
            capacity_note=capacity_note,
        )
    )


@tool
@_serialized
def record_purchase(purchases: list[dict]) -> str:
    """Add purchases to stock. The AI reads the owner's notes, confirms the
    line items with them, then calls this.

    Use the actual amount when known; otherwise place a rough figure from the
    estimated amount per purchase unit (unit_weight). Later physical counts
    always win over estimates.

    Args:
        purchases: List of purchase line items. Each element is a dict with:
            item_id: Item ID (required).
            amount: The actual amount added, in the item's registered unit.
                Provide when the real quantity is known.
            units: Number of purchase units (e.g. the 2 in "2 cones of meat").
                Without amount, the stock is estimated as units × unit_weight.
    """
    if not purchases:
        return t("purchase_empty")

    state = load_state()
    lines: list[str] = []
    for entry in purchases:
        item_id = entry.get("item_id")
        item = _find(state["items"], item_id) if item_id else None
        if item is None or not item.get("active", True):
            return _item_missing(state, "purchase_item_not_found", item_id)

        amount = entry.get("amount")
        units = entry.get("units")
        estimated = False
        if amount is None:
            if units is None:
                return t("purchase_needs_amount_or_units", name=item["name"])
            unit_weight = item.get("unit_weight")
            if unit_weight:
                # A distinct purchase unit exists (1 cone = 10kg, 1 bag = 10
                # sheets, ...) — estimate from the typical amount.
                amount = float(units) * float(unit_weight)
                estimated = True
            elif item.get("unit_type") == "count":
                amount = float(units)
            else:
                return t(
                    "purchase_no_unit_weight",
                    name=item["name"],
                    purchase_unit=item.get("purchase_unit") or t("purchase_unit_fallback"),
                )
        amount = float(amount)
        if amount < 0:
            return t("purchase_nonnegative", name=item["name"])

        before = float(item["stock"])
        after = _round_one(before + amount)
        item["stock"] = after

        # With both the actual amount and the unit count, learn the real
        # "grams per cone" from experience.
        if (
            not estimated
            and units
            and item.get("unit_type") in ("weight", "volume")
            and item.get("purchase_unit")
        ):
            samples = item.setdefault("unit_weight_samples", [])
            samples.append(_round_one(amount / float(units)))
            item["unit_weight"] = _round_one(sum(samples) / len(samples))

        state["history"].append(
            {
                "recorded_at": _now_iso(),
                "type": "purchase",
                "item_id": item["id"],
                "amount": amount,
                "units": units,
                "estimated": estimated,
            }
        )
        note = t("purchase_estimated_note") if estimated else ""
        lines.append(
            t(
                "purchase_line",
                name=item["name"],
                amount=_amount(amount, item["unit"]),
                before=_display_number(before),
                after=_amount(after, item["unit"]),
                note=note,
            )
        )

    if conflict := _save(state):
        return conflict
    return t("purchase_recorded") + "\n" + "\n".join(lines)


@tool
@_serialized
def get_stock_status() -> str:
    """Return every item's current stock, coefficient, learning phase, last
    stock count, and recent settings changes.

    Items whose book value is negative show "recount needed" instead of the
    number (principle 10).
    """
    state = load_state()
    lines: list[str] = []
    for item in state["items"]:
        if not item.get("active", True):
            continue
        ctype = _consumption_type(item)
        last_counted = item.get("last_counted")
        last_display = _fmt_dt(last_counted) if last_counted else t("never_counted")
        if ctype == "unit":
            coefficient = item.get("coefficient")
            coef_display = (
                t("status_coef_unit", coef=float(coefficient), unit=item["unit"])
                if coefficient
                else t("status_unlearned")
            )
            opened_at = item.get("opened_at_sales_count")
            opened_note = ""
            if opened_at is not None:
                used = int(round(float(item.get("sales_count", 0)) - float(opened_at)))
                opened_note = t("status_open_note", unit=item["unit"], used=used)
            lines.append(
                t(
                    "status_line_unit",
                    name=item["name"],
                    id=item["id"],
                    stock=_display_stock(item),
                    coef=coef_display,
                    growth=_growth_label(item),
                    opened_note=opened_note,
                    last=last_display,
                )
            )
        elif ctype == "count":
            lines.append(
                t(
                    "status_line_count",
                    name=item["name"],
                    id=item["id"],
                    stock=_display_stock(item),
                    last=last_display,
                )
            )
        else:
            lines.append(
                t(
                    "status_line_weight",
                    name=item["name"],
                    id=item["id"],
                    stock=_display_stock(item),
                    coef=float(item["coefficient"]),
                    growth=_growth_label(item),
                    last=last_display,
                )
            )

    settings_log = state.get("settings_log", [])
    if settings_log:
        lines.append("")
        lines.append(t("status_recent_changes"))
        for entry in settings_log[-5:]:
            lines.append(f'- {_fmt_dt(entry["changed_at"])} {entry["summary"]}')
    return _tell("\n".join(lines))


@tool
@_serialized
def get_sales_summary() -> str:
    """Return per-product average daily sales, split by venue type
    (event / solo).

    Events and solo days differ in crowd and scale, so they are averaged
    separately, never mixed.
    """
    state = load_state()
    sales = [
        entry
        for entry in state["history"]
        if entry.get("product_id") is not None and entry.get("recorded_at")
    ]
    if not sales:
        return t("sales_none")

    venue_names = {"event": t("venue_event"), "solo": t("venue_solo")}
    blocks: list[str] = []
    for venue in VENUE_TYPES:
        venue_sales = [s for s in sales if s.get("venue_type", "solo") == venue]
        if not venue_sales:
            continue
        days = {_parse_iso(s["recorded_at"]).date() for s in venue_sales}
        totals: dict[str, float] = {}
        for sale in venue_sales:
            totals[sale["product_id"]] = totals.get(sale["product_id"], 0.0) + float(
                sale["quantity"]
            )
        lines = [t("sales_block_header", venue=venue_names[venue], days=len(days))]
        for product_id, total in totals.items():
            product = _find(state["products"], product_id)
            name = product["name"] if product else product_id
            average = _round_one(total / len(days))
            lines.append(
                t(
                    "sales_product_line",
                    name=name,
                    total=_display_number(total),
                    average=_display_number(average),
                )
            )
        blocks.append("\n".join(lines))
    return "\n\n".join(blocks)


@tool
@_serialized
def get_recipes() -> str:
    """Return every product with its ID, price, and recipe (ingredient name,
    item ID, amount per serving).

    Use it to see what a product is made of, or to look up the product and
    item IDs needed by update_recipe / record_sales — never guess an ID.
    """
    state = load_state()
    if not state["products"]:
        return t("recipes_none")
    lines: list[str] = []
    for product in state["products"]:
        parts = []
        for ingredient in product["recipe"]:
            item = _find(state["items"], ingredient["item_id"])
            if item is None:
                parts.append(t("recipes_ingredient_missing", id=ingredient["item_id"]))
            else:
                parts.append(
                    t(
                        "recipes_ingredient",
                        name=item["name"],
                        id=item["id"],
                        amount=_amount(float(ingredient["qty"]), item["unit"]),
                    )
                )
        lines.append(
            t(
                "recipes_line",
                name=product["name"],
                id=product["id"],
                price=product["price"],
                recipe=_join(parts),
            )
        )
    return _tell("\n".join(lines))


@tool
@_serialized
def get_capacity() -> str:
    """Return how many more servings of each product the current stock allows,
    with the bottleneck item.

    All remaining-serving math happens in Python. Use it to judge whether
    stock will last — not to forecast sales.
    """
    state = load_state()
    if not state["products"]:
        return t("capacity_no_products")
    lines: list[str] = []
    for product in state["products"]:
        servings: float | None = None
        bottleneck = ""
        missing = ""
        unlearned: list[str] = []
        for ingredient in product["recipe"]:
            item = _find(state["items"], ingredient["item_id"])
            if item is None or not item.get("active", True):
                missing = ingredient["item_id"]
                break
            ctype = _consumption_type(item)
            qty = float(ingredient["qty"])
            if qty <= 0:
                continue
            if ctype == "unit":
                coefficient = item.get("coefficient")
                if not coefficient:
                    # An unlearned unit item can't yield a remaining count.
                    # Never silently treat it as 0 or infinity.
                    unlearned.append(item["name"])
                    continue
                available = float(item["stock"]) * float(coefficient)
                opened_at = item.get("opened_at_sales_count")
                if opened_at is not None:
                    available -= float(item.get("sales_count", 0)) - float(opened_at)
                possible = available / qty
            else:
                coefficient = float(item["coefficient"]) if ctype == "weight" else 1.0
                per_serving = qty * coefficient
                if per_serving <= 0:
                    continue
                possible = float(item["stock"]) / per_serving
            if servings is None or possible < servings:
                servings = possible
                bottleneck = item["name"]
        unlearned_note = (
            t("capacity_unlearned_note", names=_join(unlearned)) if unlearned else ""
        )
        if missing:
            lines.append(t("capacity_missing", product=product["name"], missing=missing))
        elif servings is None:
            lines.append(t("capacity_none", product=product["name"], note=unlearned_note))
        elif servings < 0:
            lines.append(
                t(
                    "capacity_negative",
                    product=product["name"],
                    bottleneck=bottleneck,
                    note=unlearned_note,
                )
            )
        else:
            lines.append(
                t(
                    "capacity_line",
                    product=product["name"],
                    servings=int(servings),
                    bottleneck=bottleneck,
                    note=unlearned_note,
                )
            )
    return "\n".join(lines)


@tool
@_serialized
def get_monthly_reconciliation(month: str = "") -> str:
    """Monthly reconciliation: per item, compare the change between physical
    counts against "purchases − recipe-basis consumption".

    Small gaps that stay inside the daily cap (±4g/serving) still show up
    when accumulated over a month (principle 12). The cause of a gap
    (missed records, waste, shrinkage) cannot be told apart, so only the
    facts are returned. Unit-tracked items are excluded — their empty-unit
    records are already a reconciliation.

    Args:
        month: Target month (YYYY-MM). Defaults to the current month.
    """
    if month:
        if not re.match(r"^\d{4}-\d{2}$", month):
            return t("month_format")
    else:
        month = _now().strftime("%Y-%m")

    state = load_state()
    findings: list[str] = []
    uncheckable: list[str] = []
    for item in state["items"]:
        if not item.get("active", True):
            continue
        ctype = _consumption_type(item)
        if ctype == "unit":
            continue
        counts = [
            entry
            for entry in item.get("count_history", [])
            if "actual_stock" in entry and entry["recorded_at"][:7] == month
        ]
        if len(counts) < 2:
            uncheckable.append(item["name"])
            continue
        first, last = counts[0], counts[-1]
        window_start = _parse_iso(first["recorded_at"])
        window_end = _parse_iso(last["recorded_at"])

        purchased = 0.0
        recipe_consumption = 0.0
        servings = 0.0
        for entry in state["history"]:
            recorded_at = entry.get("recorded_at")
            if recorded_at is None:
                continue
            at = _parse_iso(recorded_at)
            if not (window_start < at <= window_end):
                continue
            if entry.get("type") == "purchase":
                if entry.get("item_id") == item["id"]:
                    purchased += float(entry["amount"])
                continue
            product_id = entry.get("product_id")
            if product_id is None:
                continue
            product = _find(state["products"], product_id)
            if product is None:
                continue
            recipe = next(
                (i for i in product["recipe"] if i["item_id"] == item["id"]), None
            )
            if recipe:
                quantity = float(entry["quantity"])
                recipe_consumption += quantity * float(recipe["qty"])
                servings += quantity

        # Expected stock on a recipe-value basis vs. the physical count.
        # Because this bypasses the coefficient, drift that was absorbed
        # into the coefficient surfaces in the accumulation
        # (design doc, principle 12: "check monthly").
        expected_last = float(first["actual_stock"]) + purchased - recipe_consumption
        gap = _round_one(expected_last - float(last["actual_stock"]))
        bound = _bound_per_serving(item["unit"])
        allowance = _round_one(bound * servings) if ctype == "weight" else 0.0
        if gap == 0 or abs(gap) <= allowance:
            continue
        unit = item["unit"]
        period = t(
            "period_range",
            start=_fmt_dt(first["recorded_at"]),
            end=_fmt_dt(last["recorded_at"]),
        )
        direction = t("short") if gap > 0 else t("over")
        if ctype == "weight":
            findings.append(
                t(
                    "recon_weight_finding",
                    name=item["name"],
                    period=period,
                    purchased=_amount(purchased, unit),
                    consumption=_amount(_round_one(recipe_consumption), unit),
                    servings=_display_number(servings),
                    gap=_amount(abs(gap), unit),
                    direction=direction,
                    bound=_display_bound(bound),
                    unit=unit,
                    allowance=_amount(allowance, unit),
                )
            )
        else:
            findings.append(
                t(
                    "recon_count_finding",
                    name=item["name"],
                    period=period,
                    gap=_amount(abs(gap), unit),
                    direction=direction,
                )
            )

    lines = [t("recon_header", month=month)]
    if findings:
        lines.extend(findings)
    else:
        lines.append(t("recon_no_gaps"))
    if uncheckable:
        lines.append(t("recon_uncheckable", names=_join(uncheckable)))
    return _tell("\n".join(lines))


# ---------------------------------------------------------------------------
# Structure changes (passphrase required; change history kept)
# ---------------------------------------------------------------------------


def _recorded_months(state: dict[str, Any]) -> list[str]:
    """Every YYYY-MM that holds a sale, a purchase, or a stock count."""
    months = {
        entry["recorded_at"][:7]
        for entry in state["history"]
        if entry.get("recorded_at")
    }
    for item in state["items"]:
        months.update(
            entry["recorded_at"][:7]
            for entry in item.get("count_history", [])
            if entry.get("recorded_at")
        )
    return sorted(months)


def _date_parts(recorded_at: str) -> tuple[str, str]:
    """Split an ISO timestamp into a date and a time column. Spreadsheets
    sort and filter those; one long ISO string they do not."""
    return recorded_at[:10], recorded_at[11:19]


def _write_csv(path: Path, header: list[str], rows: list[list[Any]]) -> None:
    """Write one sheet. utf-8-sig, because Excel reads a BOM-less UTF-8 CSV
    as the local codepage and turns every non-ASCII name into mojibake."""
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerows(rows)


def _export_sales(state: dict[str, Any], month: str) -> tuple[list[str], list[list[Any]]]:
    header = [
        t("csv_date"), t("csv_time"), t("csv_venue"), t("csv_product_id"),
        t("csv_product"), t("csv_quantity"), t("csv_unit_price"), t("csv_revenue"),
    ]
    rows: list[list[Any]] = []
    for entry in state["history"]:
        recorded_at = entry.get("recorded_at")
        if entry.get("product_id") is None or not recorded_at:
            continue
        if month and recorded_at[:7] != month:
            continue
        product = _find(state["products"], entry["product_id"])
        price = float(product["price"]) if product else 0.0
        quantity = float(entry.get("quantity") or 0)
        date, clock = _date_parts(recorded_at)
        rows.append([
            date, clock, entry.get("venue_type", "solo"), entry["product_id"],
            product["name"] if product else entry["product_id"],
            _display_number(quantity), _display_number(price),
            _display_number(price * quantity),
        ])
    return header, rows


def _export_purchases(state: dict[str, Any], month: str) -> tuple[list[str], list[list[Any]]]:
    header = [
        t("csv_date"), t("csv_time"), t("csv_item_id"), t("csv_item"),
        t("csv_amount"), t("csv_unit"), t("csv_units"), t("csv_purchase_unit"),
        t("csv_estimated"),
    ]
    rows: list[list[Any]] = []
    for entry in state["history"]:
        recorded_at = entry.get("recorded_at")
        if entry.get("type") != "purchase" or not recorded_at:
            continue
        if month and recorded_at[:7] != month:
            continue
        item = _find(state["items"], entry.get("item_id"))
        date, clock = _date_parts(recorded_at)
        rows.append([
            date, clock, entry.get("item_id"),
            item["name"] if item else entry.get("item_id"),
            _display_number(float(entry.get("amount") or 0)),
            item["unit"] if item else "",
            _display_number(float(entry["units"])) if entry.get("units") else "",
            (item.get("purchase_unit") or "") if item else "",
            t("csv_yes") if entry.get("estimated") else t("csv_no"),
        ])
    return header, rows


def _export_counts(state: dict[str, Any], month: str) -> tuple[list[str], list[list[Any]]]:
    header = [
        t("csv_date"), t("csv_time"), t("csv_item_id"), t("csv_item"),
        t("csv_actual_stock"), t("csv_book_stock"), t("csv_gap"), t("csv_unit"),
        t("csv_coef_before"), t("csv_coef_after"),
    ]
    rows: list[list[Any]] = []
    for item in state["items"]:
        for entry in item.get("count_history", []):
            recorded_at = entry.get("recorded_at")
            # Only real physical counts. count_history also holds empty-unit
            # and unit-change events, which have no measured stock.
            if "actual_stock" not in entry or not recorded_at:
                continue
            if month and recorded_at[:7] != month:
                continue
            actual = float(entry["actual_stock"])
            book = float(entry.get("calculated_stock") or 0.0)
            date, clock = _date_parts(recorded_at)
            rows.append([
                date, clock, item["id"], item["name"],
                _display_number(actual), _display_number(book),
                _display_number(_round_one(actual - book)), item["unit"],
                entry.get("coefficient_before", ""), entry.get("coefficient_after", ""),
            ])
    rows.sort(key=lambda row: (row[0], row[1]))
    return header, rows


def _export_stock(state: dict[str, Any]) -> tuple[list[str], list[list[Any]]]:
    header = [
        t("csv_item_id"), t("csv_item"), t("csv_stock"), t("csv_unit"),
        t("csv_consumption_type"), t("csv_coefficient"), t("csv_last_counted"),
    ]
    rows: list[list[Any]] = []
    for item in state["items"]:
        if not item.get("active", True):
            continue
        coefficient = item.get("coefficient")
        rows.append([
            item["id"], item["name"], _display_number(float(item["stock"])),
            item["unit"], _consumption_type(item),
            _display_number(float(coefficient)) if coefficient else "",
            (item.get("last_counted") or "")[:10],
        ])
    return header, rows


def _export_recipes(state: dict[str, Any]) -> tuple[list[str], list[list[Any]]]:
    header = [
        t("csv_product_id"), t("csv_product"), t("csv_price"), t("csv_item_id"),
        t("csv_item"), t("csv_qty_per_serving"), t("csv_unit"),
    ]
    rows: list[list[Any]] = []
    for product in state["products"]:
        for ingredient in product["recipe"]:
            item = _find(state["items"], ingredient["item_id"])
            rows.append([
                product["id"], product["name"], _display_number(float(product["price"])),
                ingredient["item_id"], item["name"] if item else ingredient["item_id"],
                _display_number(float(ingredient["qty"])), item["unit"] if item else "",
            ])
    return header, rows


# Columns that hold a quantity rather than a label. CSV writes everything as
# text and lets the spreadsheet guess; a workbook can do better, so these
# become real numbers and a SUM or a pivot over them works straight away.
_NUMERIC_COLUMNS = (
    "csv_quantity", "csv_unit_price", "csv_revenue", "csv_amount", "csv_units",
    "csv_actual_stock", "csv_book_stock", "csv_gap", "csv_coef_before",
    "csv_coef_after", "csv_stock", "csv_coefficient", "csv_price",
    "csv_qty_per_serving",
)


def _as_number(value: Any) -> Any:
    """A numeric column's cell as a number, or unchanged if it is blank or
    not a number after all (a coefficient column can hold either)."""
    if isinstance(value, (int, float)):
        return value
    text = str(value).strip()
    if not text:
        return ""
    try:
        number = float(text)
    except ValueError:
        return value
    return int(number) if number.is_integer() else number


def _export_sheets(
    state: dict[str, Any], month: str
) -> list[tuple[str, str, list[str], list[list[Any]]]]:
    """The five sheets both exporters write, as (filename, label, header, rows).

    Sales, purchases and counts honour ``month``; stock and recipes are the
    current state whatever the month, because a photograph of last August's
    shelf does not exist.
    """
    return [
        ("sales.csv", t("export_label_sales"), *_export_sales(state, month)),
        ("purchases.csv", t("export_label_purchases"), *_export_purchases(state, month)),
        ("stock_counts.csv", t("export_label_counts"), *_export_counts(state, month)),
        ("stock_now.csv", t("export_label_stock"), *_export_stock(state)),
        ("recipes.csv", t("export_label_recipes"), *_export_recipes(state)),
    ]


def _nothing_to_export(
    state: dict[str, Any],
    month: str,
    sheets: list[tuple[str, str, list[str], list[list[Any]]]],
) -> str:
    """A sentence to return when there is no point writing anything, else "".

    stock_now and recipes always have rows, so "is there anything here" has to
    be asked of the three record sheets. Otherwise a typo'd month writes a
    bundle that looks complete and contains no records.
    """
    if sum(len(rows) for _, _, _, rows in sheets[:3]):
        return ""
    if not month:
        return t("export_none")
    # Name the months that exist. The model has no idea what year it is, so
    # "August" becomes a guess; without this it guesses again.
    return (
        f'{t("export_no_records_for_month", month=month)}\n'
        f'{t("export_months_available", months=_join(_recorded_months(state)))}'
    )


def _write_xlsx(
    path: Path, sheets: list[tuple[str, str, list[str], list[list[Any]]]]
) -> None:
    """Write every sheet into one workbook: bold frozen header, dates as dates,
    quantities as numbers, columns wide enough to read without dragging."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    numeric_labels = {t(key) for key in _NUMERIC_COLUMNS}
    date_labels = {t("csv_date"), t("csv_last_counted")}

    book = Workbook()
    book.remove(book.active)
    for _, label, header, rows in sheets:
        sheet = book.create_sheet(label[:31])
        sheet.append(header)
        numeric = [index for index, name in enumerate(header) if name in numeric_labels]
        dates = [index for index, name in enumerate(header) if name in date_labels]

        for row in rows:
            cells = list(row)
            for index in numeric:
                cells[index] = _as_number(cells[index])
            for index in dates:
                try:
                    cells[index] = date.fromisoformat(str(cells[index]))
                except ValueError:
                    pass  # an unparseable date stays the text it was
            sheet.append(cells)
            for index in dates:
                # openpyxl would otherwise show a date as its serial number.
                cell = sheet.cell(row=sheet.max_row, column=index + 1)
                cell.number_format = "yyyy-mm-dd"

        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        if rows:
            sheet.auto_filter.ref = sheet.dimensions
        for index, name in enumerate(header, start=1):
            widest = max(
                [_cell_width(name)] + [_cell_width(row[index - 1]) for row in rows]
            )
            sheet.column_dimensions[get_column_letter(index)].width = min(widest + 2, 40)

    book.save(path)


def _cell_width(value: Any) -> int:
    """Rough display width. A Japanese character occupies about two columns of
    a spreadsheet's default font, so counting characters underestimates it."""
    text = str(value)
    return sum(1 if character.isascii() else 2 for character in text)


@tool
@_serialized
def export_excel(month: str = "") -> str:
    """Write the shop's records to one Excel workbook (.xlsx) — the file to
    hand to an accountant, or to keep as a copy outside teruo.

    One file, five sheets: sales, purchases, stock counts, stock right now,
    recipes. Excel opens it, and so does Google Sheets — the five sheets
    become five tabs of one document. Quantities are written as numbers with
    the unit in its own column, so the spreadsheet totals them. Purchase costs are not recorded
    anywhere in teruo, so the purchases sheet carries quantities only, never
    money. Use this unless the owner specifically asks for CSV.

    Args:
        month: Limit sales, purchases, and stock counts to one month
            (YYYY-MM). Empty writes every record. The stock and recipe sheets
            are always the current state, whatever the month.
    """
    if month and not re.match(r"^\d{4}-\d{2}$", month):
        return t("month_format")

    state = load_state()
    sheets = _export_sheets(state, month)
    empty = _nothing_to_export(state, month, sheets)
    if empty:
        return empty

    period = month or "all"
    folder = STATE_PATH.parent / "exports" / period
    path = folder / f"teruo-{period}.xlsx"
    try:
        folder.mkdir(parents=True, exist_ok=True)
        _write_xlsx(path, sheets)
    except ImportError:
        return t("excel_unavailable")
    except OSError as error:
        return t("export_failed", reason=error)

    _offer([path])
    lines = [
        t(
            "excel_done",
            count=len(sheets),
            path=path,
            period=month or t("export_period_all"),
        )
    ]
    lines.extend(
        t("excel_line", label=label, rows=len(rows))
        for _, label, _, rows in sheets
    )
    lines.append(t("excel_open_hint"))
    return _tell("\n".join(lines))


@tool
@_serialized
def export_csv(month: str = "") -> str:
    """Write the shop's records as five CSV files, for feeding another system:
    accounting software, a script, anything that wants plain text.

    Prefer export_excel for a spreadsheet. It writes the same records as one
    .xlsx, and Google Sheets opens that as five tabs of one document — five
    CSVs would become five separate documents there. Use this tool when the
    owner asks for CSV by name, or names a system that takes CSV.

    Five files: sales, purchases, stock counts, stock right now, recipes.
    Numbers are written bare with the unit in its own column, so the
    spreadsheet can total them. Purchase costs are not recorded anywhere in
    teruo, so the purchase file carries quantities only, never money.

    Args:
        month: Limit sales, purchases, and stock counts to one month
            (YYYY-MM). Empty writes every record. The stock and recipe files
            are always the current state, whatever the month.
    """
    if month and not re.match(r"^\d{4}-\d{2}$", month):
        return t("month_format")

    state = load_state()
    sheets = _export_sheets(state, month)
    empty = _nothing_to_export(state, month, sheets)
    if empty:
        return empty

    folder = STATE_PATH.parent / "exports" / (month or "all")
    try:
        folder.mkdir(parents=True, exist_ok=True)
        for filename, _, header, rows in sheets:
            _write_csv(folder / filename, header, rows)
    except OSError as error:
        return t("export_failed", reason=error)

    _offer([folder / filename for filename, _, _, _ in sheets])
    lines = [
        t(
            "export_done",
            count=len(sheets),
            folder=folder,
            period=month or t("export_period_all"),
        )
    ]
    lines.extend(
        t("export_line", file=filename, label=label, rows=len(rows))
        for filename, label, _, rows in sheets
    )
    lines.append(t("export_open_hint"))
    return _tell("\n".join(lines))


@tool
@_serialized
def register_item(
    item_id: str,
    name: str,
    unit: str,
    unit_type: str,
    stock: float,
    tier: int = 1,
    consumption_type: str = "",
    purchase_unit: str = "",
    unit_weight: float = 0.0,
    passphrase: str = "",
) -> str:
    """Add an item (ingredient or supply). A structure change — passphrase required.

    Weight-item coefficients always start at 1.0 (never let the owner pick
    one); physical counts grow them. Unit items start with no value — the
    first empty unit (record_unit_used) settles it.

    Args:
        item_id: Lowercase snake_case ID, e.g. sauce_yogurt. Named by the AI.
        name: Display name, e.g. Yogurt sauce.
        unit: Consumption unit, e.g. g, pc, sheet, bottle, tub.
        unit_type: "weight" / "volume" / "count".
        stock: Current stock. A rough figure is fine.
        tier: Management tier. A provisional guess is fine.
        consumption_type: How consumption is counted. "count" (decreases
            one-for-one) / "weight" (portioned by hand; coefficient learned
            from stock counts) / "unit" (servings per bottle/piece/tub;
            learned from empty units).
            Decided by the onboarding question "what do you count as one
            when you use it?" — never guessed from the item name. If omitted,
            falls back to unit_type (count→count, others→weight).
        purchase_unit: Only when the purchase unit differs from the
            consumption unit, e.g. cone, bag.
        unit_weight: Estimated amount per purchase unit. Only with
            purchase_unit.
        passphrase: The owner's passphrase.
    """
    state = load_state()
    if error := _check_passphrase(state, passphrase):
        return error
    if not ITEM_ID_PATTERN.match(item_id):
        return t("item_id_format")
    if _find(state["items"], item_id) is not None:
        return t("item_exists", item_id=item_id)
    if unit_type not in UNIT_TYPES:
        return t("unit_type_invalid")
    if consumption_type and consumption_type not in CONSUMPTION_TYPES:
        return t("consumption_type_invalid")
    if stock < 0:
        return t("stock_nonnegative")

    if not consumption_type:
        consumption_type = "count" if unit_type == "count" else "weight"
    item: dict[str, Any] = {
        "id": item_id,
        "name": name,
        "tier": tier,
        "unit": unit,
        "unit_type": unit_type,
        "consumption_type": consumption_type,
        "stock": _round_one(stock),
        "coefficient": 1.0,
        "last_counted": None,
        "count_history": [],
        "active": True,
    }
    if consumption_type == "unit":
        item["coefficient"] = None
        item["sales_count"] = 0
        item["opened_at_sales_count"] = None
        item["unit_history"] = []
    if purchase_unit:
        item["purchase_unit"] = purchase_unit
        if unit_weight > 0:
            item["unit_weight"] = unit_weight
    state["items"].append(item)
    _log_setting(
        state,
        "register_item",
        t("log_item_added", name=name, stock=_amount(float(stock), unit)),
        None,
        {k: item[k] for k in ("id", "name", "unit", "consumption_type", "stock")},
    )
    if conflict := _save(state):
        return conflict
    stock_display = _amount(float(stock), unit)
    if consumption_type == "unit":
        return t("item_registered_unit", name=name, stock=stock_display, unit=unit)
    if consumption_type == "count":
        return t("item_registered_count", name=name, stock=stock_display)
    return t("item_registered_weight", name=name, stock=stock_display)


@tool
@_serialized
def register_product(
    product_id: str,
    name: str,
    price: int,
    recipe: list[dict],
    passphrase: str = "",
) -> str:
    """Add a product (menu entry). A structure change — passphrase required.

    If the recipe mentions unregistered items, register them with
    register_item first, then call this (keeping the owner unaware of the
    ordering is the AI's job).

    Args:
        product_id: Lowercase snake_case ID, e.g. kebab_sand.
        name: Display name, e.g. Kebab sandwich.
        price: Price including tax (yen).
        recipe: Ingredient list. Each element is a dict of
            {"item_id": item ID, "qty": amount per serving}.
        passphrase: The owner's passphrase.
    """
    state = load_state()
    if error := _check_passphrase(state, passphrase):
        return error
    if not ITEM_ID_PATTERN.match(product_id):
        return t("product_id_format")
    if _find(state["products"], product_id) is not None:
        return t("product_exists", product_id=product_id)
    if price < 0:
        return t("price_nonnegative")
    if not recipe:
        return t("recipe_empty")

    seen: set[str] = set()
    missing: list[str] = []
    for ingredient in recipe:
        item_id = ingredient.get("item_id")
        qty = ingredient.get("qty")
        if not item_id or qty is None or float(qty) <= 0:
            return t("recipe_element_invalid")
        if item_id in seen:
            return t("recipe_duplicate", item_id=item_id)
        seen.add(item_id)
        item = _find(state["items"], item_id)
        if item is None or not item.get("active", True):
            missing.append(item_id)
    if missing:
        return t("recipe_missing_items", names=_join(missing))

    product = {
        "id": product_id,
        "name": name,
        "price": price,
        "recipe": [
            {"item_id": i["item_id"], "qty": float(i["qty"])} for i in recipe
        ],
    }
    state["products"].append(product)
    recipe_text = _join(
        f'{_find(state["items"], i["item_id"])["name"]} '
        f'{_amount(float(i["qty"]), _find(state["items"], i["item_id"])["unit"])}'
        for i in product["recipe"]
    )
    _log_setting(
        state,
        "register_product",
        t("log_product_added", name=name, price=price, recipe=recipe_text),
        None,
        product,
    )
    if conflict := _save(state):
        return conflict
    return t("product_registered", name=name, price=price, recipe=recipe_text)


@tool
@_serialized
def update_recipe(
    product_id: str,
    item_id: str,
    qty: float,
    passphrase: str = "",
) -> str:
    """Change, add, or remove one ingredient in a product recipe.
    A structure change — passphrase required.

    Args:
        product_id: Product ID.
        item_id: Item ID.
        qty: New amount per serving. 0 removes the ingredient from the recipe.
        passphrase: The owner's passphrase.
    """
    state = load_state()
    if error := _check_passphrase(state, passphrase):
        return error
    product = _find(state["products"], product_id)
    if product is None:
        return _product_missing(state, product_id)
    item = _find(state["items"], item_id)
    if item is None or not item.get("active", True):
        return _item_missing(state, "item_unregistered", item_id)
    if qty < 0:
        return t("recipe_qty_nonnegative")

    existing = next(
        (i for i in product["recipe"] if i["item_id"] == item_id), None
    )
    unit = item["unit"]
    if qty == 0:
        if existing is None:
            return t("recipe_not_include", product=product["name"], item=item["name"])
        product["recipe"] = [
            i for i in product["recipe"] if i["item_id"] != item_id
        ]
        summary = t(
            "log_recipe_removed",
            product=product["name"],
            item=item["name"],
            amount=_amount(float(existing["qty"]), unit),
        )
        before, after = float(existing["qty"]), None
        message = t("recipe_removed", product=product["name"], item=item["name"])
    elif existing is None:
        product["recipe"].append({"item_id": item_id, "qty": float(qty)})
        summary = t(
            "log_recipe_added",
            product=product["name"],
            item=item["name"],
            amount=_amount(float(qty), unit),
        )
        before, after = None, float(qty)
        message = t(
            "recipe_added",
            product=product["name"],
            item=item["name"],
            amount=_amount(float(qty), unit),
        )
    else:
        before = float(existing["qty"])
        existing["qty"] = float(qty)
        summary = t(
            "log_recipe_changed",
            product=product["name"],
            item=item["name"],
            before=_amount(before, unit),
            after=_amount(float(qty), unit),
        )
        after = float(qty)
        message = t(
            "recipe_changed",
            product=product["name"],
            item=item["name"],
            before=_amount(before, unit),
            after=_amount(float(qty), unit),
        )
    _log_setting(state, "update_recipe", summary, before, after)
    if conflict := _save(state):
        return conflict
    return message


@tool
@_serialized
def update_item(
    item_id: str,
    new_name: str = "",
    new_unit: str = "",
    new_stock: float = -1.0,
    new_purchase_unit: str = "",
    new_unit_weight: float = 0.0,
    passphrase: str = "",
) -> str:
    """Fix an item's name, unit, or purchase unit. A structure change —
    passphrase required.

    Unit-change rules:
    - Within the same kind (g→kg etc.): stock, recipes, and the estimated
      purchase amount convert automatically.
    - Across kinds (g→pc etc.): no conversion is possible, so a recounted
      stock (new_stock) is required. The coefficient resets to 1.0 and the
      learning phase restarts (past records are kept).

    Args:
        item_id: Item ID.
        new_name: New display name (only when changing).
        new_unit: New consumption unit (only when changing), e.g. kg, pc.
        new_stock: On a cross-kind unit change, the recounted stock in the
            new unit.
        new_purchase_unit: Purchase unit (set or change it once known),
            e.g. cone, bag.
        new_unit_weight: Estimated amount per purchase unit, in the
            consumption unit. E.g. 10000 for a 10kg cone.
        passphrase: The owner's passphrase.
    """
    state = load_state()
    if error := _check_passphrase(state, passphrase):
        return error
    item = _find(state["items"], item_id)
    if item is None or not item.get("active", True):
        return _item_missing(state, "item_not_found", item_id)
    if not new_name and not new_unit and not new_purchase_unit and new_unit_weight <= 0:
        return t("update_item_nothing")

    messages: list[str] = []

    if new_purchase_unit or new_unit_weight > 0:
        if not new_purchase_unit and not item.get("purchase_unit"):
            return t("unit_weight_alone")
        before_purchase = {
            "purchase_unit": item.get("purchase_unit"),
            "unit_weight": item.get("unit_weight"),
        }
        if new_purchase_unit:
            item["purchase_unit"] = new_purchase_unit
        if new_unit_weight > 0:
            item["unit_weight"] = new_unit_weight
        purchase_unit = item["purchase_unit"]
        weight_note = (
            t(
                "purchase_weight_note",
                purchase_unit=purchase_unit,
                amount=_amount(float(item["unit_weight"]), item["unit"]),
            )
            if item.get("unit_weight")
            else ""
        )
        _log_setting(
            state,
            "update_item",
            t(
                "log_purchase_unit_set",
                name=item["name"],
                purchase_unit=purchase_unit,
                note=weight_note,
            ),
            before_purchase,
            {
                "purchase_unit": item.get("purchase_unit"),
                "unit_weight": item.get("unit_weight"),
            },
        )
        messages.append(
            t("purchase_unit_registered", purchase_unit=purchase_unit, note=weight_note)
        )

    if new_name and new_name != item["name"]:
        old_name = item["name"]
        item["name"] = new_name
        _log_setting(
            state,
            "update_item",
            t("log_item_renamed", old=old_name, new=new_name),
            old_name,
            new_name,
        )
        messages.append(t("item_renamed", old=old_name, new=new_name))

    if new_unit and new_unit != item["unit"]:
        old_unit = item["unit"]
        old_stock = float(item["stock"])
        old_coefficient = float(item["coefficient"])
        factor = _conversion_factor(old_unit, new_unit)
        before_snapshot = {
            "unit": old_unit,
            "unit_type": item["unit_type"],
            "stock": old_stock,
            "coefficient": old_coefficient,
        }

        if factor is not None:
            # Same kind: mechanically convert stock, recipes, and the
            # estimate. The coefficient keeps its meaning, so it stays.
            item["unit"] = new_unit
            item["stock"] = _round_one(old_stock * factor)
            if item.get("unit_weight"):
                item["unit_weight"] = float(item["unit_weight"]) * factor
            affected: list[str] = []
            for product in state["products"]:
                for ingredient in product["recipe"]:
                    if ingredient["item_id"] == item_id:
                        ingredient["qty"] = float(ingredient["qty"]) * factor
                        affected.append(product["name"])
            _log_setting(
                state,
                "update_item",
                t(
                    "log_unit_converted",
                    name=item["name"],
                    old_unit=old_unit,
                    new_unit=new_unit,
                    old_stock=_amount(old_stock, old_unit),
                    new_stock=_amount(float(item["stock"]), new_unit),
                ),
                before_snapshot,
                {"unit": new_unit, "stock": item["stock"]},
            )
            note = (
                t("recipes_converted_note", names=_join(dict.fromkeys(affected)))
                if affected
                else ""
            )
            messages.append(
                t(
                    "unit_converted",
                    old_unit=old_unit,
                    new_unit=new_unit,
                    stock=_amount(float(item["stock"]), new_unit),
                    note=note,
                )
            )
        else:
            # Different kind: no conversion. A recount is mandatory
            # (design doc §2-4).
            new_kind = _unit_kind(new_unit)
            if new_stock < 0:
                return t(
                    "unit_change_needs_recount",
                    name=item["name"],
                    old_unit=old_unit,
                    new_unit=new_unit,
                    old_stock=_amount(old_stock, old_unit),
                )
            item["unit"] = new_unit
            if new_kind:
                item["unit_type"] = new_kind
            item["stock"] = _round_one(new_stock)
            item["coefficient"] = 1.0
            item["last_counted"] = None
            item.setdefault("count_history", []).append(
                {
                    "recorded_at": _now_iso(),
                    "event": "unit_change",
                    "before": before_snapshot,
                    "after": {
                        "unit": new_unit,
                        "unit_type": item["unit_type"],
                        "stock": item["stock"],
                        "coefficient": 1.0,
                    },
                }
            )
            affected_products = [
                product["name"]
                for product in state["products"]
                if any(i["item_id"] == item_id for i in product["recipe"])
            ]
            _log_setting(
                state,
                "update_item",
                t(
                    "log_unit_changed",
                    name=item["name"],
                    old_unit=old_unit,
                    new_unit=new_unit,
                    old_stock=_amount(old_stock, old_unit),
                    new_stock=_amount(float(new_stock), new_unit),
                ),
                before_snapshot,
                {
                    "unit": new_unit,
                    "unit_type": item["unit_type"],
                    "stock": item["stock"],
                    "coefficient": 1.0,
                },
            )
            warning = (
                t("recipes_old_unit_warning", names=_join(affected_products))
                if affected_products
                else ""
            )
            messages.append(
                t(
                    "unit_changed",
                    old_unit=old_unit,
                    new_unit=new_unit,
                    stock=_amount(float(new_stock), new_unit),
                    warning=warning,
                )
            )
    elif new_unit:
        messages.append(t("unit_already", unit=new_unit))

    if conflict := _save(state):
        return conflict
    return " ".join(messages) if messages else t("nothing_changed")


@tool
@_serialized
def delete_product(
    product_id: str,
    confirm: bool = False,
    passphrase: str = "",
) -> str:
    """Delete a product. Confirmation required; a structure change, so the
    passphrase is required too.

    Deleting a product keeps its items (ingredients) and sales history.
    Call with confirm=True only after the owner has clearly confirmed the
    deletion.

    Args:
        product_id: Product ID.
        confirm: True only once the owner has approved the deletion.
        passphrase: The owner's passphrase.
    """
    state = load_state()
    if error := _check_passphrase(state, passphrase):
        return error
    product = _find(state["products"], product_id)
    if product is None:
        return _product_missing(state, product_id)
    if not confirm:
        return t("delete_confirm", name=product["name"], price=product["price"])
    state["products"] = [p for p in state["products"] if p["id"] != product_id]
    _log_setting(
        state,
        "delete_product",
        t("log_product_deleted", name=product["name"], price=product["price"]),
        product,
        None,
    )
    if conflict := _save(state):
        return conflict
    return t("product_deleted", name=product["name"])


@tool
@_serialized
def reset_shop(
    confirm: bool = False,
    passphrase: str = "",
) -> str:
    """Set this shop aside and start over with an empty one.

    Call it first with no arguments: the reply says what would be set aside
    and whether the passphrase is needed. Call again with confirm=True only
    after the owner has clearly said yes. Nothing is deleted — the current
    state file is renamed with a timestamp, and the onboarding interview
    starts by itself afterwards.

    Args:
        confirm: True only once the owner has approved the reset.
        passphrase: The owner's passphrase (required when one is set).
    """
    state = load_state()
    if error := _check_passphrase(state, passphrase):
        return error
    items = len(state.get("items") or [])
    products = len(state.get("products") or [])
    history = len(state.get("history") or [])
    if not confirm:
        return t("reset_confirm", items=items, products=products, history=history)
    archive = archive_and_reset(_now().strftime("%Y%m%d-%H%M%S"))
    return _tell(t("reset_done", archive=archive.name))


@tool
@_serialized
def update_config(
    passphrase: str = "",
    new_shop_name: str = "",
    new_passphrase: str = "",
    new_notify_email: str = "",
) -> str:
    """Set the shop's name, the passphrase, and the email address notified of
    settings changes.

    On first setup (nothing stored yet) these can be set directly. Changing
    them later requires the current passphrase.

    Args:
        passphrase: The current passphrase (only needed when changing).
        new_shop_name: What to call this shop — the shop's name, or the
            owner's. teruo greets them by it at every startup.
        new_passphrase: The new passphrase.
        new_notify_email: Email address to notify about settings changes.
    """
    if not new_shop_name and not new_passphrase and not new_notify_email:
        return t("config_nothing")
    state = load_state()
    if error := _check_passphrase(state, passphrase):
        return error
    config = state.setdefault("config", {"passphrase": None, "notify_email": None})
    messages: list[str] = []
    if new_shop_name:
        shop = new_shop_name.strip()
        old_shop = config.get("shop_name")
        config["shop_name"] = shop
        _log_setting(
            state, "update_config", t("log_shop_name_set", shop=shop), old_shop, shop
        )
        messages.append(t("shop_name_set", shop=shop))
    if new_passphrase:
        config["passphrase"] = new_passphrase
        _log_setting(state, "update_config", t("log_passphrase_set"), None, t("hidden"))
        messages.append(t("passphrase_set"))
    if new_notify_email:
        old_email = config.get("notify_email")
        config["notify_email"] = new_notify_email
        _log_setting(
            state,
            "update_config",
            t("log_email_set", email=new_notify_email),
            old_email,
            new_notify_email,
        )
        messages.append(t("email_set", email=new_notify_email))
    if conflict := _save(state):
        return conflict
    return " ".join(messages)
